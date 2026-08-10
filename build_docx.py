import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FIGDIR = "paper/figs_png"
FIGS = {
    "FIG_TOY": ("fig_toygraph.png", "Fig. 1.  A toy precinct dual graph. Precincts are nodes, colored by district; an edge joins two adjacent precincts, and edges that cross a district boundary (red) are the cut edges the detector reads.", 2.5),
    "FIG_ARCH": ("fig1_architecture.png", "Fig. 2.  Model architecture. A GIN student is distilled against a frozen random teacher over the neutral ensemble; the pooled node-and-graph discrepancy gives a plan score s(P) that flags an enacted plan and localizes planted edits, run as a topology-only and a +demographic variant.", 1.35),
    "FIG_LOCAL": ("fig2_localization.png", "Fig. 3.  Node-localization AUC of the learned score (blue) against a boundary and cut-edge baseline given the same ensemble conditioning (red), by state and planting condition.", 2.45),
    "FIG_CUTEDGE": ("fig3_cutedges.png", "Fig. 4.  Cut edges of each enacted plan (red line) against the distribution over its neutral ensemble (blue).", 2.5),
}
TABLES = {
    "TABLE_I": ("TABLE I.  Outlier Verdict of Each Enacted Plan",
        ["Detector", "NC (R)", "PA (R)", "MD (D)"], [
        ("mean–median", "99.1", "64.9", "67.5"),
        ("efficiency gap", "100", "98.7", "100"),
        ("declination", "100", "95.7", "100"),
        ("cut edges", "53.7", "100", "100"),
        ("GNN score", "82.0", "99.6", "99.2")],
        "Extreme percentile of each enacted plan against one neutral ensemble (2016 presidential vote); 99 is the flag threshold. Rows 1–3 are vote-based, rows 4–5 geometry-based.",
        [Inches(1.45), Inches(0.63), Inches(0.63), Inches(0.63)]),
    "TABLE_II": ("TABLE II.  Multi-Election Robustness",
        ["State", "mean–med.", "eff. gap", "declin.", "lean", "geom. p"], [
        ("NC (R)", "5/5", "4/5", "4/5", "0.49", "0.36"),
        ("PA (R)", "4/9", "3/9", "4/9", "0.52", "0.009"),
        ("MD (D)", "0/6", "2/6", "4/6", "0.55", "0.016")],
        "Fraction of statewide races under which each vote statistic flags the enacted plan beyond the 99th percentile. “lean” is the mean two-party Democratic share; “geom. p” is the topology-only detector's election-invariant outlier p-value.",
        [Inches(0.55), Inches(0.68), Inches(0.55), Inches(0.55), Inches(0.45), Inches(0.55)]),
}

def set_cols(section, n, space=0.25):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sectPr.append(cols)
    cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(int(space * 1440)))

def runs_from(p, text, size=10):
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = p.add_run(seg); r.font.size = Pt(size); r.font.name = "Times New Roman"; r.bold = (i % 2 == 1)

def fixed_layout(t, widths):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tblW = OxmlElement("w:tblW"); tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(w.inches for w in widths) * 1440))); tblPr.append(tblW)
    grid = t._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i].inches * 1440)))
    for row in t.rows:
        for j, wd in enumerate(widths):
            row.cells[j].width = wd

def add_table(doc, spec):
    cap_txt, header, rows, note = spec[0], spec[1], spec[2], spec[3]
    widths = spec[4] if len(spec) > 4 else None
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(cap_txt); cr.font.size = Pt(8); cr.font.name = "Times New Roman"; cr.font.small_caps = True
    cp.paragraph_format.space_before = Pt(4); cp.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows) + 1, cols=len(header)); t.style = "Table Grid"; t.autofit = False
    if widths is None:
        widths = [Inches(2.4), Inches(0.95)] if len(header) == 2 else [Inches(1.95), Inches(0.55), Inches(0.75)]
    fixed_layout(t, widths)
    fs = 7 if len(header) >= 5 else 8
    for j, htext in enumerate(header):
        c = t.cell(0, j); c.paragraphs[0].text = ""
        rr = c.paragraphs[0].add_run(htext); rr.bold = True; rr.font.size = Pt(fs); rr.font.name = "Times New Roman"
    for i, r in enumerate(rows):
        for j, val in enumerate(r):
            c = t.cell(i + 1, j); c.paragraphs[0].text = ""
            rr = c.paragraphs[0].add_run(str(val)); rr.font.size = Pt(fs); rr.font.name = "Times New Roman"
    nrows = len(t.rows)
    for ri, row in enumerate(t.rows):
        tr = row._tr; trPr = tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = OxmlElement("w:trPr"); tr.insert(0, trPr)
        trPr.append(OxmlElement("w:cantSplit"))
        if ri < nrows - 1 or note:
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
    if note:
        npp = doc.add_paragraph(); nr = npp.add_run(note); nr.font.size = Pt(6); nr.italic = True; nr.font.name = "Times New Roman"
        npp.paragraph_format.space_after = Pt(4)

def add_fig(doc, marker):
    fname, cap, w = FIGS[marker]
    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(2); fp.paragraph_format.keep_with_next = True
    fp.add_run().add_picture(f"{FIGDIR}/{fname}", width=Inches(w))
    cpp = doc.add_paragraph(); cr = cpp.add_run(cap); cr.font.size = Pt(8); cr.font.name = "Times New Roman"
    cpp.paragraph_format.space_after = Pt(2)

doc = Document()
doc.core_properties.author = "Jacob Crainic"
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0
sec0 = doc.sections[0]
sec0.top_margin = Inches(0.75); sec0.bottom_margin = Inches(1.0)
sec0.left_margin = Inches(0.62); sec0.right_margin = Inches(0.62)
sec0.page_height = Inches(11); sec0.page_width = Inches(8.5)

blocks = [b.strip() for b in open("paper/urtc_paper.md").read().split("\n\n") if b.strip()]
set_cols(sec0, 1)
title = blocks[0].lstrip("# ").strip()
doc.core_properties.title = title
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(title); r.bold = False; r.font.size = Pt(24); r.font.name = "Times New Roman"
p.paragraph_format.space_after = Pt(4)
for line, sz in [("Jacob Crainic", 11), ("Philosophy, Politics, Economics, and Law", 10),
                 ("University of Florida", 10), ("Gainesville, FL, USA", 10), ("jacobcrainic2008@gmail.com", 10)]:
    ap = doc.add_paragraph(); ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ar = ap.add_run(line); ar.font.size = Pt(sz); ar.font.name = "Times New Roman"
ap.paragraph_format.space_after = Pt(3)

newsec = doc.add_section(WD_SECTION.CONTINUOUS)
newsec.top_margin = Inches(0.75); newsec.bottom_margin = Inches(1.0)
newsec.left_margin = Inches(0.62); newsec.right_margin = Inches(0.62)
set_cols(newsec, 2)

for b in blocks[1:]:
    flat = re.sub(r"\s+", " ", b).strip()
    m = re.match(r"^\[\[(\w+)\]\]$", b.strip())
    if m:
        key = m.group(1)
        if key in FIGS:
            add_fig(doc, key)
        elif key in TABLES:
            add_table(doc, TABLES[key])
        continue
    if b.startswith("### "):
        hp = doc.add_paragraph(); hp.paragraph_format.space_before = Pt(5); hp.paragraph_format.space_after = Pt(1)
        hr = hp.add_run(b[4:].strip()); hr.italic = True; hr.font.size = Pt(10); hr.font.name = "Times New Roman"
        continue
    if b.startswith("## "):
        hp = doc.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_before = Pt(6); hp.paragraph_format.space_after = Pt(2)
        hr = hp.add_run(b[3:].strip()); hr.font.small_caps = True; hr.font.size = Pt(10); hr.font.name = "Times New Roman"; hp.paragraph_format.keep_with_next = True
        continue
    if flat.startswith("**Abstract**"):
        ap = doc.add_paragraph(); ap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body = re.sub(r"\*\*Abstract\*\*\s*[—-]?", "", flat).strip()
        r0 = ap.add_run("Abstract—"); r0.bold = True; r0.italic = True; r0.font.size = Pt(9); r0.font.name = "Times New Roman"
        r1 = ap.add_run(body); r1.bold = True; r1.font.size = Pt(9); r1.font.name = "Times New Roman"
        continue
    if flat.startswith("**Keywords**"):
        kp = doc.add_paragraph(); body = re.sub(r"\*\*Keywords\*\*\s*[—-]?", "", flat).strip()
        r0 = kp.add_run("Keywords—"); r0.italic = True; r0.font.size = Pt(9); r0.font.name = "Times New Roman"
        r1 = kp.add_run(body); r1.italic = True; r1.font.size = Pt(9); r1.font.name = "Times New Roman"
        kp.paragraph_format.space_after = Pt(6)
        continue
    if re.match(r"^\[\d+\]", flat):
        rp = doc.add_paragraph(); rp.paragraph_format.space_after = Pt(0); rp.paragraph_format.widow_control = False
        for i, seg in enumerate(flat.split("*")):
            if not seg:
                continue
            rr = rp.add_run(seg); rr.font.size = Pt(8); rr.font.name = "Times New Roman"; rr.italic = (i % 2 == 1)
        continue
    bp = doc.add_paragraph(); bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bp.paragraph_format.first_line_indent = Inches(0.2)
    runs_from(bp, flat, size=10)

doc.save("paper/urtc_paper.docx")
d2 = Document("paper/urtc_paper.docx")
print("wrote paper/urtc_paper.docx | sections", len(d2.sections), "| paras", len(d2.paragraphs), "| tables", len(d2.tables))
